import os

from PIL import Image, ImageDraw
from pptx.enum.shapes import PP_PLACEHOLDER

from app.document_diff.config import (
    PPTX_EXACT_TEXT_SIMILARITY,
    PPTX_FOOTER_MAX_TEXT_LENGTH,
    PPTX_FOOTER_START_RATIO,
    PPTX_HIGHLIGHT_BORDER_WIDTH,
    PPTX_HIGHLIGHT_COLORS,
    PPTX_HIGHLIGHT_PADDING,
    PPTX_MIN_SHAPE_MATCH_SCORE,
    PPTX_PARTIAL_MATCH_MAX_DISTANCE,
    PPTX_PARTIAL_TEXT_SIMILARITY,
    PPTX_POSITION_TOLERANCE,
    PPTX_SHAPE_DISTANCE_NORMALIZER,
    PPTX_SHAPE_DISTANCE_WEIGHT,
    PPTX_SHAPE_TEXT_WEIGHT,
)
from app.document_diff.pptx_alignment import normalize_text, text_similarity


def _is_footer(shape, height):
    placeholders = (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.HEADER, PP_PLACEHOLDER.DATE)
    if shape.is_placeholder and shape.placeholder_format.type in placeholders:
        return True
    return (
        shape.has_text_frame
        and shape.top.pt > height * PPTX_FOOTER_START_RATIO
        and len(shape.text.strip()) < PPTX_FOOTER_MAX_TEXT_LENGTH
    )


def shapes(slide, height):
    result = []
    for index, shape in enumerate(slide.shapes):
        if _is_footer(shape, height) or (not shape.has_text_frame and not shape.shape_type):
            continue
        result.append(
            {
                'index': index,
                'text': shape.text.strip() if shape.has_text_frame else '',
                'left': shape.left.pt,
                'top': shape.top.pt,
                'width': shape.width.pt,
                'height': shape.height.pt,
            }
        )
    return result


def match_shapes(shapes_a, shapes_b):
    matched, remaining_a, remaining_b = [], list(shapes_a), list(shapes_b)
    for shape_a in list(remaining_a):
        candidates = []
        for shape_b in remaining_b:
            distance = ((shape_a['left'] - shape_b['left']) ** 2 + (shape_a['top'] - shape_b['top']) ** 2) ** 0.5
            similarity = text_similarity(shape_a['text'], shape_b['text'])
            score = PPTX_SHAPE_TEXT_WEIGHT * similarity + PPTX_SHAPE_DISTANCE_WEIGHT * max(
                0, 1 - distance / PPTX_SHAPE_DISTANCE_NORMALIZER
            )
            if similarity >= PPTX_EXACT_TEXT_SIMILARITY or (
                similarity > PPTX_PARTIAL_TEXT_SIMILARITY and distance < PPTX_PARTIAL_MATCH_MAX_DISTANCE
            ):
                candidates.append((score, shape_b))
        if candidates and max(candidates, key=lambda item: item[0])[0] > PPTX_MIN_SHAPE_MATCH_SCORE:
            _, shape_b = max(candidates, key=lambda item: item[0])
            matched.append((shape_a, shape_b))
            remaining_a.remove(shape_a)
            remaining_b.remove(shape_b)
    return matched, remaining_a, remaining_b


def draw_highlights(path, highlights, width, height):
    if not highlights or not os.path.exists(path):
        return
    with Image.open(path) as source:
        image = source.convert('RGBA')
    overlay = Image.new('RGBA', image.size, (255, 255, 255, 0))
    drawing = ImageDraw.Draw(overlay)
    for item in highlights:
        x0 = item['left'] / width * image.width - PPTX_HIGHLIGHT_PADDING
        y0 = item['top'] / height * image.height - PPTX_HIGHLIGHT_PADDING
        x1 = (item['left'] + item['width']) / width * image.width + PPTX_HIGHLIGHT_PADDING
        y1 = (item['top'] + item['height']) / height * image.height + PPTX_HIGHLIGHT_PADDING
        color = PPTX_HIGHLIGHT_COLORS[item['type']]
        drawing.rectangle(
            (x0, y0, x1, y1), fill=color['fill'], outline=color['border'], width=PPTX_HIGHLIGHT_BORDER_WIDTH
        )
    Image.alpha_composite(image, overlay).convert('RGB').save(path)
    image.close()
    overlay.close()


def _position_changed(shape_a, shape_b, dimensions_a, dimensions_b):
    width_a, height_a = dimensions_a
    width_b, height_b = dimensions_b
    return (
        abs(shape_a['left'] / width_a - shape_b['left'] / width_b) > PPTX_POSITION_TOLERANCE
        or abs(shape_a['top'] / height_a - shape_b['top'] / height_b) > PPTX_POSITION_TOLERANCE
    )


def slide_changes(slide_a, slide_b, dimensions_a, dimensions_b):
    matched, deleted, added = match_shapes(shapes(slide_a, dimensions_a[1]), shapes(slide_b, dimensions_b[1]))
    highlights_a = [{**shape, 'type': 'deleted'} for shape in deleted]
    highlights_b = [{**shape, 'type': 'added'} for shape in added]
    types = {'deleted'} if deleted else set()
    types.update({'added'} if added else set())
    for shape_a, shape_b in matched:
        if _position_changed(shape_a, shape_b, dimensions_a, dimensions_b):
            highlights_a.append({**shape_a, 'type': 'position'})
            highlights_b.append({**shape_b, 'type': 'position'})
            types.add('shift')
        if normalize_text(shape_a['text']) != normalize_text(shape_b['text']):
            highlights_a.append({**shape_a, 'type': 'modified'})
            highlights_b.append({**shape_b, 'type': 'modified'})
            types.add('modified')
    return highlights_a, highlights_b, types
