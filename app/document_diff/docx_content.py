import copy
import difflib
import shutil

import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph

DELETED_COLOR = RGBColor.from_string('DC3545')
INSERTED_COLOR = RGBColor.from_string('00C800')
DELETED_FILL = 'F8D7DA'
INSERTED_FILL = 'D6F5D6'

def _copy_run_format(source, target):
    source_properties = source._r.find(qn('w:rPr'))
    if source_properties is None:
        return
    target_properties = copy.deepcopy(source_properties)
    existing = target._r.find(qn('w:rPr'))
    if existing is not None:
        target._r.remove(existing)
    target._r.insert(0, target_properties)


def _mark_run(run, change):
    fill = DELETED_FILL if change == 'delete' else INSERTED_FILL
    run.font.color.rgb = DELETED_COLOR if change == 'delete' else INSERTED_COLOR
    run.font.strike = change == 'delete'
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), fill)
    run._r.get_or_add_rPr().append(shading)


def _words_with_runs(paragraph):
    words, word, source_run = [], '', None
    for run in paragraph.runs:
        for character in run.text:
            if character.isspace():
                if word:
                    words.append((word, source_run))
                    word, source_run = '', None
            else:
                source_run = source_run or run
                word += character
    if word:
        words.append((word, source_run))
    return words


def _insert_after(body, reference, element):
    if reference is not None:
        reference.addnext(element)
    elif len(body):
        body[0].addprevious(element)
    else:
        body.append(element)


def _mark_entire_paragraph(paragraph, change):
    for run in paragraph.runs:
        _mark_run(run, change)


def _add_changed_pair(body, deleted, inserted, document):
    deleted_words = _words_with_runs(deleted)
    inserted_words = _words_with_runs(inserted)
    deleted_texts = [word for word, _ in deleted_words]
    inserted_texts = [word for word, _ in inserted_words]
    opcodes = difflib.SequenceMatcher(None, deleted_texts, inserted_texts).get_opcodes()

    combined = []
    for tag, start_a, end_a, start_b, end_b in opcodes:
        if tag == 'equal':
            combined.extend((word, run, None) for word, run in deleted_words[start_a:end_a])
            continue
        if tag in ('delete', 'replace'):
            combined.extend((word, run, 'delete') for word, run in deleted_words[start_a:end_a])
        if tag in ('insert', 'replace'):
            combined.extend((word, run, 'insert') for word, run in inserted_words[start_b:end_b])

    for run in list(deleted.runs):
        deleted._p.remove(run._r)
    for index, (word, source_run, change) in enumerate(combined):
        run = deleted.add_run(word + (' ' if index < len(combined) - 1 else ''))
        if source_run is not None:
            _copy_run_format(source_run, run)
        if change is not None:
            _mark_run(run, change)
    return deleted._p


def build_combined_doc(source_path, paragraphs_b, opcodes, output_path):
    shutil.copy2(source_path, output_path)
    document = docx.Document(output_path)
    body = document.element.body
    output_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
    for tag, start_a, end_a, start_b, end_b in opcodes:
        if tag == 'equal':
            continue
        deleted, inserted = output_paragraphs[start_a:end_a], paragraphs_b[start_b:end_b]
        reference = output_paragraphs[start_a - 1]._p if start_a else None
        for paragraph_a, paragraph_b in zip(deleted, inserted):
            reference = _add_changed_pair(body, paragraph_a, paragraph_b, document)
        for paragraph in deleted[len(inserted):]:
            _mark_entire_paragraph(paragraph, 'delete')
            reference = paragraph._p
        for paragraph in inserted[len(deleted):]:
            element = copy.deepcopy(paragraph._p)
            _insert_after(body, reference, element)
            copied = Paragraph(element, document)
            _mark_entire_paragraph(copied, 'insert')
            reference = element
    document.save(output_path)