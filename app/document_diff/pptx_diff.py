import os

import pymupdf as fitz
from PIL import Image
from pptx import Presentation

from app.document_diff.pptx_alignment import align_slides, dhash
from app.document_diff.pptx_render import draw_highlights, slide_changes
from app.utils.converter import convert_to

PAGE_WIDTH = 1600
COLUMN_WIDTH = 750
PAGE_MARGIN = 50
IMAGE_MARGIN = 40
MIN_PAGE_HEIGHT = 520


def _slide_path(folder, comparison_id, side, index):
    return os.path.join(folder, f'{comparison_id}_{side}_slide_{index}.png')


def _render_slides(pdf, folder, comparison_id, side):
    slides = []
    for index, page in enumerate(pdf):
        path = _slide_path(folder, comparison_id, side, index)
        page.get_pixmap(dpi=120).save(path)
        with Image.open(path) as image:
            slides.append({'index': index, 'text': page.get_text(), 'dhash': dhash(image)})
    return slides


def _render_pdfs(path_a, path_b, folder, comparison_id):
    pdf_a, pdf_b = fitz.open(os.path.splitext(path_a)[0] + '.pdf'), fitz.open(os.path.splitext(path_b)[0] + '.pdf')
    try:
        return _render_slides(pdf_a, folder, comparison_id, 'a'), _render_slides(pdf_b, folder, comparison_id, 'b')
    finally:
        pdf_a.close()
        pdf_b.close()


def _mark_unmatched(item, presentations, dimensions, folder, comparison_id, stats):
    status, side = item['status'], 'a' if item['status'] == 'deleted' else 'b'
    index = item[f'slide_{side}']
    presentation = presentations[0 if side == 'a' else 1]
    width, height = dimensions[0 if side == 'a' else 1]
    draw_highlights(_slide_path(folder, comparison_id, side, index), [{'left': 0, 'top': 0, 'width': width, 'height': height, 'type': status}], width, height)
    item['status_label'] = status
    stats[status] += 1


def _mark_matched(item, presentations, dimensions, folder, comparison_id, stats):
    index_a, index_b = item['slide_a'], item['slide_b']
    highlights_a, highlights_b, types = slide_changes(
        presentations[0].slides[index_a], presentations[1].slides[index_b], dimensions[0], dimensions[1]
    )
    draw_highlights(_slide_path(folder, comparison_id, 'a', index_a), highlights_a, *dimensions[0])
    draw_highlights(_slide_path(folder, comparison_id, 'b', index_b), highlights_b, *dimensions[1])
    if not types:
        item['status_label'] = 'unchanged'
        stats['unchanged'] += 1
    elif types & {'modified', 'deleted', 'added'}:
        item['status_label'] = 'modified'
        stats['modified'] += 1
    else:
        item['status_label'] = 'shift'
        stats['shift'] += 1


def _save_preview_pdf(mapping, folder, comparison_id):
    output = fitz.open()
    try:
        for item in mapping:
            paths = [
                _slide_path(folder, comparison_id, side, item[f'slide_{side}']) if item[f'slide_{side}'] is not None else None
                for side in ('a', 'b')
            ]
            existing = [path for path in paths if path and os.path.exists(path)]
            if not existing:
                continue
            column_heights = []
            for path in paths:
                if path and os.path.exists(path):
                    with Image.open(path) as image:
                        width, height = image.size
                    column_heights.append(int(COLUMN_WIDTH * height / width))
                else:
                    column_heights.append(0)
            page_column_height = max(column_heights)
            page = output.new_page(width=PAGE_WIDTH, height=max(MIN_PAGE_HEIGHT, page_column_height + 2 * IMAGE_MARGIN))
            offsets = (PAGE_MARGIN, PAGE_MARGIN + COLUMN_WIDTH + PAGE_MARGIN)
            for offset, path, column_height in zip(offsets, paths, column_heights):
                if path and os.path.exists(path):
                    page.insert_image(
                        fitz.Rect(offset, IMAGE_MARGIN, offset + COLUMN_WIDTH, IMAGE_MARGIN + column_height),
                        filename=path,
                    )
        output.save(os.path.join(folder, f'{comparison_id}_diff_highlighted.pdf'))
    finally:
        output.close()


def compare_pptx(path_a, path_b, comparison_id):
    if convert_to(path_a) is None or convert_to(path_b) is None:
        raise ValueError('Не удалось преобразовать презентацию в PDF')
    folder = os.path.dirname(path_a)
    slides_a, slides_b = _render_pdfs(path_a, path_b, folder, comparison_id)
    presentations = Presentation(path_a), Presentation(path_b)
    dimensions = [(presentation.slide_width.pt, presentation.slide_height.pt) for presentation in presentations]
    mapping = align_slides(slides_a, slides_b)
    stats = dict.fromkeys(('modified', 'added', 'deleted', 'shift', 'unchanged'), 0)
    for item in mapping:
        if item['status'] in ('deleted', 'added'):
            _mark_unmatched(item, presentations, dimensions, folder, comparison_id, stats)
        else:
            _mark_matched(item, presentations, dimensions, folder, comparison_id, stats)
    _save_preview_pdf(mapping, folder, comparison_id)
    return {'mapping': mapping, 'stats': stats}