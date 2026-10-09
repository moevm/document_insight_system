import difflib
import os

import docx
import pymupdf as fitz

from app.document_diff.docx_content import build_combined_doc
from app.document_diff.docx_pdf import collect_fragments, save_pdf


def compare_docx(docx_a_path, docx_b_path, comparison_id):
    document_a, document_b = docx.Document(docx_a_path), docx.Document(docx_b_path)
    paragraphs_a = [paragraph for paragraph in document_a.paragraphs if paragraph.text.strip()]
    paragraphs_b = [paragraph for paragraph in document_b.paragraphs if paragraph.text.strip()]
    text_a = [' '.join(paragraph.text.split()) for paragraph in paragraphs_a]
    text_b = [' '.join(paragraph.text.split()) for paragraph in paragraphs_b]
    opcodes = difflib.SequenceMatcher(None, text_a, text_b).get_opcodes()
    folder = os.path.dirname(docx_a_path)
    output_path = os.path.join(folder, f'{comparison_id}_diff.docx')
    build_combined_doc(docx_a_path, paragraphs_b, opcodes, output_path)
    pdf_a_path = os.path.join(folder, f'{comparison_id}_1_hl.pdf')
    pdf_b_path = os.path.join(folder, f'{comparison_id}_2_hl.pdf')
    pdf_a, pdf_b = fitz.open(pdf_a_path), fitz.open(pdf_b_path)
    try:
        fragments = collect_fragments(pdf_a, pdf_b, paragraphs_a, paragraphs_b, opcodes)
        save_pdf(pdf_a, pdf_a_path)
        save_pdf(pdf_b, pdf_b_path)
    except Exception:
        pdf_a.close()
        pdf_b.close()
        raise
    return {'fragments': fragments}
