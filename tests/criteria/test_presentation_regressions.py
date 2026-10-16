import base64
from io import BytesIO
from types import SimpleNamespace

import pytest
import requests
from pptx import Presentation
from pptx.util import Inches

from app.main.checks.presentation_checks.aspect_ratio_check import PresAspectRatioCheck
from app.main.checks.presentation_checks.banned_words import PresBannedWordsCheck
from app.main.checks.presentation_checks.empty_slide_check import PresEmptySlideCheck
from app.main.checks.presentation_checks.find_theme_in_pres import FindThemeInPres
from app.main.checks.presentation_checks.image_share import PresImageShareCheck
from app.main.checks.presentation_checks.search_keyword import SearchKeyWord
from app.main.checks.presentation_checks.sld_enum import SldEnumCheck
from app.main.checks.presentation_checks.verify_git_link import PresVerifyGitLinkCheck
from app.main.presentations.pptx.presentation_pptx import PresentationPPTX


def file_info(titles=None, texts=None):
    file = SimpleNamespace(get_titles=lambda: titles or [], get_text_from_slides=lambda: texts or [], slides=[])
    return {'file': file, 'filename': 'slides.pptx', 'pdf_id': 'test', 'file_type': {'type': 'pres'}}


def test_nonstandard_aspect_ratio_returns_failure():
    info = file_info()
    info['file'].prs = Presentation()
    info['file'].prs.slide_width = Inches(10)
    info['file'].prs.slide_height = Inches(10)
    assert PresAspectRatioCheck(info).check()['score'] == 0
    info['file'].prs.slide_height = Inches(7.5)
    assert PresAspectRatioCheck(info).check()['score'] == 1


def test_git_failure_not_overwritten_and_no_order_dependency(monkeypatch):
    seen = []

    def get(url, **kwargs):
        seen.append((url, kwargs))
        return SimpleNamespace(status_code=404, content=b'Not found')

    monkeypatch.setattr(requests, 'get', get)
    info = file_info(['Апробация'], ['github.com/owner/repo'])
    result = PresVerifyGitLinkCheck(info).check()
    assert result['score'] == 0
    assert seen == [('https://github.com/owner/repo', {'timeout': 10})]


def test_git_empty_and_missing_link(monkeypatch):
    monkeypatch.setattr(
        requests,
        'get',
        lambda *a, **k: SimpleNamespace(status_code=200, content=b'<html>This repository is empty.</html>'),
    )
    assert (
        PresVerifyGitLinkCheck(file_info(['Апробация'], ['https://github.com/owner/repo']), deep_check=True).check()[
            'score'
        ]
        == 0
    )
    assert PresVerifyGitLinkCheck(file_info(['Апробация'], ['Репозиторий (github)'])).check()['score'] == 0
    assert PresVerifyGitLinkCheck(file_info(['Раздел'], ['Просто текст'])).check()['score'] == 1


def test_git_timeout_is_reported(monkeypatch):
    def timeout(*args, **kwargs):
        raise requests.exceptions.Timeout()

    monkeypatch.setattr(requests, 'get', timeout)
    assert PresVerifyGitLinkCheck(file_info(['Раздел'], ['https://github.com/owner/repo'])).check()['score'] == 0


def test_topic_comes_from_title_slide():
    info = file_info(['Анализ данных', 'Цель и задачи', 'Анализ'], ['Анализ данных', 'Цель и задачи', 'Анализ'])
    assert FindThemeInPres(info, limit=40).check()['score'] == 1
    info = file_info(
        ['Анализ данных', 'Цель и задачи', 'Результаты'], ['Анализ данных', 'Цель и задачи', 'Анализ данных']
    )
    assert FindThemeInPres(info).check()['score'] == 0  # Body is not a heading.
    assert FindThemeInPres(file_info()).check()['score'] == 0
    assert FindThemeInPres(file_info(['', 'Раздел'])).check()['score'] == 0


def test_image_only_pptx_parses_and_is_not_empty(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    image = base64.b64decode(
        'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jK1kAAAAASUVORK5CYII='
    )
    slide.shapes.add_picture(BytesIO(image), Inches(1), Inches(1))
    path = tmp_path / 'image-only.PPTX'
    prs.save(path)
    info = file_info()
    info['file'] = PresentationPPTX(str(path))
    assert PresEmptySlideCheck(info).check()['score'] == 1


@pytest.mark.parametrize('criterion', [PresImageShareCheck, SldEnumCheck])
def test_empty_presentation_does_not_divide_by_zero(criterion):
    assert criterion(file_info()).check()['score'] == 0


def test_banned_words_count_repeated_occurrences_and_phrases():
    info = file_info(['Раздел'], ['мы мы мы мы'])
    result = PresBannedWordsCheck(info, words=['мы']).check()
    assert result['score'] == 0.5
    assert 'Слайд №1' in result['verdict'][0]
    info = file_info(['Раздел'], ['Интуитивно понятный интерфейс'])
    assert PresBannedWordsCheck(info, words=['интуитивно понятный'], min_count=0, max_count=0).check()['score'] == 0


def test_keyword_string_is_not_split_into_letters():
    assert SearchKeyWord(file_info(['Раздел'], ['Другой текст']), 'Актуальность').check()['score'] == 0
    assert SearchKeyWord(file_info(['Раздел'], ['Актуальность']), 'Актуальность').check()['score'] == 1
