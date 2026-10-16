from types import SimpleNamespace

import docx
import pytest

from app.main.check_packs.base_criterion_pack import BaseCriterionPack
from app.main.checks.check_abbreviations import is_abbreviation_explained
from app.main.checks.check_analyzer import AVAILABLE_CHECKS
from app.main.checks.report_checks.headers_at_page_top_check import ReportHeadersAtPageTopCheck
from app.main.checks.report_checks.image_references import ImageReferences
from app.main.checks.report_checks.spelling_check import SpellingCheck
from app.main.checks.report_checks.sw_section_lit_reference import SWSectionLiteratureReferenceCheck
from app.main.checks.report_checks.table_references import TableReferences
from app.main.reports.pdf_document.pdf_document_manager import PdfDocumentManager
from app.utils.decimal_places_check import DecimalPlacesCheck


def test_pdf_page_slice_preserves_absolute_numbers(make_pdf):
    path = make_pdf(['one', 'two', 'three'])
    pdf = PdfDocumentManager(str(path), str(path))
    assert list(pdf.get_text_on_page(start_page=2, end_page=3)) == [2, 3]
    assert 'two' in pdf.get_text_on_page(start_page=2, end_page=2)[2]
    assert pdf.get_text_on_page(end_page=0) == {}


def test_registry_does_not_expose_helper_criteria():
    assert None not in AVAILABLE_CHECKS['report']
    assert len(AVAILABLE_CHECKS['report']) == 39
    assert len(AVAILABLE_CHECKS['pres']) == 23


def test_pack_initialization_failure_is_visible_and_fails(md_document):
    info = md_document('# Название\n\n## Раздел\n\nТекст')
    pack = BaseCriterionPack([['simple_check'], ['missing_check']], info['file_type'])
    pack.init(info)
    results, score, passed = pack.check()
    assert not passed and score == 0
    assert [result['id'] for result in results] == ['simple_check', 'missing_check']
    pack = BaseCriterionPack(
        [['simple_check'], ['simple_check'], ['sw_sections_size_check', {'sections_info': 'BAD'}]], info['file_type']
    )
    pack.init(info)
    results, score, passed = pack.check()
    assert not passed and len(results) == 2
    assert results[1]['score'] == 0


def test_caption_references_include_last_body_paragraph(md_document):
    info = md_document('# Название\n\n## Раздел\n\n![Рисунок 1 — Пример](image.png)\n\nСм. рисунок 1.')
    assert ImageReferences(info).check()['score'] == 1
    info = md_document('# Название\n\n## Раздел\n\nТаблица 1 — Данные\n\n| A |\n|---|\n| 1 |\n\nСм. таблицу 1.')
    assert TableReferences(info).check()['score'] == 1
    info = md_document('# Название\n\n## Раздел\n\n| A |\n|---|\n| 1 |')
    assert TableReferences(info).check()['score'] == 0


def test_duplicate_caption_numbers_fail(md_document):
    info = md_document('# Название\n\n## Раздел\n\nТаблица 1 — А\n\nТаблица 1 — Б\n\nСм. таблицу 1.')
    assert TableReferences(info).check()['score'] == 0


def test_headers_in_middle_of_page_are_failure(docx_document):
    document = docx.Document()
    document.add_heading('Введение', 2)
    document.add_paragraph('Текст')
    info = docx_document(document)
    info['file'].first_lines = ['номер Введение', '', '', '']
    info['file'].page_count = 4
    assert ReportHeadersAtPageTopCheck(info).check()['score'] == 0


def test_spelling_uses_md_first_page_and_correct_absolute_keys(md_document, monkeypatch):
    info = md_document('# Название\n\n## Раздел\n\nТекст')
    seen = []
    report_types = []

    def find_literature_page(work_type):
        report_types.append(work_type)
        return 0

    monkeypatch.setattr(info['file'], 'find_literature_page', find_literature_page)
    wrapper = SimpleNamespace(error_counter=0, check_pages=lambda pages: seen.extend(pages), stats=lambda: '')
    monkeypatch.setattr('app.main.checks.report_checks.spelling_check.LanguageToolWrapper', lambda: wrapper)
    assert SpellingCheck(info).check()['score'] == 1
    assert [page for page, text in seen] == [1]
    assert report_types == ['VKR']


def test_sw_literature_does_not_double_count_across_subchapters(md_document):
    info = md_document(
        '# Название\n\n## Принцип отбора аналогов\n\nИсточник [1].\n\n### Подраздел\n\nТот же источник [1].'
    )
    criterion = SWSectionLiteratureReferenceCheck(info, sections_info={'Принцип отбора аналогов': {'count': 2}})
    assert criterion.check()['score'] == 0


def test_abbreviation_explanation_must_be_at_first_occurrence():
    assert not is_abbreviation_explained('АБВ', 'АБВ используется.\nАБВ — Анализ Больших Выборок')
    assert is_abbreviation_explained('АБВ', 'АБВ — Анализ Больших Выборок')


@pytest.mark.parametrize('number, expected', [('1234.5678', 4), ('12,3456', 4), ('1,234.5678', 4)])
def test_decimal_long_integer_part(number, expected):
    checker = DecimalPlacesCheck({'file_type': {'type': 'report'}}, max_decimal_places=2)
    violations = checker.find_violations_in_text(number)
    assert len(violations) == 1
    assert violations[0][1] == expected


def test_effective_style_keeps_direct_overrides_and_reports_zero_values():
    from docx.shared import Cm

    from app.main.reports.docx_uploader.style import Style

    document = docx.Document()
    document.styles['Normal'].paragraph_format.first_line_indent = Cm(1.25)
    document.styles['Normal'].paragraph_format.line_spacing = 1.5
    paragraph = document.add_paragraph('Текст')
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.runs[0]
    run.bold = False
    actual = Style(run, paragraph)
    expected = Style()
    expected.bold = True
    expected.first_line_indent_cm = 1.25
    expected.line_spacing = 1.5
    errors = []
    assert not actual.matches(expected, errors)
    assert len(errors) == 3
    assert actual.first_line_indent_cm == 0
    assert actual.line_spacing == 1


def test_docx_start_alignment_parses(docx_document):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = docx.Document()
    paragraph = document.add_paragraph('Текст')
    alignment = OxmlElement('w:jc')
    alignment.set(qn('w:val'), 'start')
    paragraph._p.get_or_add_pPr().append(alignment)
    info = docx_document(document)
    assert info['file'].styled_paragraphs[0]['runs'][0]['style'].alignment == 0


def test_word_content_control_paragraph_is_visible(docx_document):
    from docx.oxml import OxmlElement

    document = docx.Document()
    paragraph = document.add_paragraph('Видимый текст поля')
    control = OxmlElement('w:sdt')
    content = OxmlElement('w:sdtContent')
    document.element.body.insert(0, control)
    content.append(paragraph._p)
    control.append(content)
    info = docx_document(document)
    assert info['file'].styled_paragraphs[0]['text'] == 'Видимый текст поля'


@pytest.mark.parametrize('entry', [[], 'simple_check', ['simple_check', []], ['simple_check', {}, 'extra']])
def test_invalid_pack_entry_never_silently_passes(md_document, entry):
    info = md_document('# Название\n\n## Раздел\n\nТекст')
    pack = BaseCriterionPack([['simple_check'], entry], info['file_type'])
    pack.init(info)
    results, score, passed = pack.check()
    assert not passed and score == 0
    assert results[-1]['id'] == 'invalid_criterion'


def test_equivalent_pack_entries_are_counted_once(md_document):
    info = md_document('# Название\n\n## Раздел\n\nТекст')
    pack = BaseCriterionPack([['simple_check'], ['simple_check', {}]], info['file_type'])
    pack.init(info)
    results, score, passed = pack.check()
    assert passed and score == 1
    assert len(results) == 1
