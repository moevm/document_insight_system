import docx
import pytest

from app.main.checks.report_checks.abbreviations_check import ReportAbbreviationsCheck
from app.main.checks.report_checks.banned_words_in_literature import BannedWordsInLiteratureCheck
from app.main.checks.report_checks.chapters import ReportChapters
from app.main.checks.report_checks.compare_goal_and_content import CompareGoalAndContentCheck
from app.main.checks.report_checks.compare_tasks_and_content import CompareTasksAndContentCheck
from app.main.checks.report_checks.decimal_places import ReportDecimalPlacesCheck
from app.main.checks.report_checks.empty_task_page_check import EmptyTaskPageCheck
from app.main.checks.report_checks.first_pages_check import ReportFirstPagesCheck
from app.main.checks.report_checks.headers_references import ReportHeadersReferencesCheck
from app.main.checks.report_checks.literature_references import ReferencesToLiteratureCheck
from app.main.checks.report_checks.main_character_check import ReportMainCharacterCheck
from app.main.checks.report_checks.max_abstract_size_check import ReportMaxSizeOfAbstractCheck
from app.main.checks.report_checks.paragraphs_count_check import ReportParagraphsCountCheck
from app.main.checks.report_checks.section_component import ReportSectionComponent
from app.main.checks.report_checks.short_sections_check import ReportShortSectionsCheck
from app.main.checks.report_checks.sw_keywords_check import SWKeywordsCheck
from app.main.checks.report_checks.sw_section_banned_words import SWSectionBannedWordsCheck
from app.main.checks.report_checks.sw_section_lit_reference import SWSectionLiteratureReferenceCheck
from app.main.checks.report_checks.sw_section_size import SWSectionSizeCheck
from app.main.checks.report_checks.sw_tasks import SWTasksCheck


def test_md_counts_real_pages_and_preserves_bold(md_document):
    info = md_document('# Название\n\n## Введение\n\n**Цель** и **задачи** определены.')
    file = info['file']
    assert file.page_counter() == file.pdf_file.page_count_all == 1
    assert file.pdf_filepath.endswith('assignment.version.pdf')
    assert file.make_chapters()[0]['style'] == 'heading 1'
    assert file.make_chapters()[0]['text'] == 'Название'
    result = ReportSectionComponent(info, patterns=['цель', 'задачи']).check()
    assert result['score'] == 1, result
    assert (
        ReportSectionComponent(
            md_document('# Название\n\n## Введение\n\nЦель и задачи определены.'), patterns=['цель', 'задачи']
        ).check()['score']
        == 0
    )


@pytest.mark.parametrize('file_format', ['md', 'docx'])
def test_component_can_disable_bold_without_disabling_presence(md_document, docx_document, file_format):
    if file_format == 'md':
        info = md_document('# Название\n\n## Задачи\n\nВыполнить обзор и сравнение.')
    else:
        document = docx.Document()
        document.add_heading('Задачи', 2)
        document.add_paragraph('Выполнить обзор и сравнение.')
        info = docx_document(document)
    options = {'chapter': 'Задачи', 'patterns': ['обзор', 'сравнен']}
    assert ReportSectionComponent(info, **options).check()['score'] == 0
    assert ReportSectionComponent(info, **options, check_bold=False).check()['score'] == 1
    options['patterns'].append('критер')
    assert ReportSectionComponent(info, **options, check_bold=False).check()['score'] == 0


def test_component_can_span_docx_runs(docx_document):
    document = docx.Document()
    document.add_heading('Введение', 2)
    paragraph = document.add_paragraph()
    for text in ('Це', 'ль', ' и задачи'):
        paragraph.add_run(text).bold = True
    info = docx_document(document)
    assert ReportSectionComponent(info, patterns=['цель', 'задачи']).check()['score'] == 1


def test_md_parser_tables_trailing_caption_and_idempotency(md_document):
    file = md_document(
        '# Название\n\n## Раздел\n\nТаблица 1 — Данные\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\nТаблица 2 — Без таблицы'
    )['file']
    assert len(file.tables) == 1
    count = len(file.styled_paragraphs)
    file.parse_effective_styles()
    assert len(file.styled_paragraphs) == count
    assert file.styled_paragraphs[-1]['style_name'] == 'вкр_подпись таблицы'
    file.make_chapters()
    chapters = file.chapters_to_str()
    tree = file.build_chapter_tree(chapters)
    assert tree[0]['children'][0]['name'] == 'Раздел'
    assert all('node' in chapter for chapter in chapters)


@pytest.mark.parametrize(
    'criterion', [SWSectionSizeCheck, SWSectionLiteratureReferenceCheck, SWSectionBannedWordsCheck]
)
def test_sw_missing_sections_fail_and_defaults_initialize(md_document, criterion):
    info = md_document('# Название\n\n## Другой раздел\n\nТекст.')
    assert criterion(info).check()['score'] == 0
    with pytest.raises(ValueError):
        criterion(info, sections_info='WRONG')


def test_sw_custom_dictionary_not_mutated(md_document):
    info = md_document('# Название\n\n## Проблема\n\nОтсутствует решение.')
    settings = {'Проблема': ['отсутствовать']}
    assert SWSectionBannedWordsCheck(info, sections_info=settings).check()['score'] == 0
    assert settings == {'Проблема': ['отсутствовать']}
    size = SWSectionSizeCheck(info, sections_info={'Проблема': {'words': 1, 'sentences': 1}})
    assert size.check()['score'] == 0


def test_sw_tasks_missing_and_bounds(md_document):
    assert SWTasksCheck(md_document('# Название\n\n## Введение\n\nТекст.')).check()['score'] == 0
    info = md_document('# Название\n\n## Задачи\n\n1. Первая\n2. Вторая\n3. Третья')
    assert SWTasksCheck(info).check()['score'] == 1
    assert SWTasksCheck(info, max_tasks=2).check()['score'] == 0


def test_sw_keywords_without_nltk_data(md_document):
    info = md_document('# Название\n\nКлючевые слова: система, анализ, данные\n\n## Раздел\n\nАнализ данных системы.')
    assert SWKeywordsCheck(info).check()['score'] == 1
    missing = md_document('# Название\n\nКлючевые слова: система, анализ, робот\n\n## Раздел\n\nАнализ системы.')
    assert SWKeywordsCheck(missing).check()['score'] == 0


def test_literature_explicit_limits_and_repeated_checks(md_document):
    info = md_document(
        '# Название\n\n## Раздел\n\nИсточник [1].\n\n## Список использованных источников\n\n1. Первый источник'
    )
    criterion = ReferencesToLiteratureCheck(info, min_ref=2)
    assert criterion.check()['score'] == 0
    criterion = ReferencesToLiteratureCheck(info, min_ref=1, max_ref=1)
    assert criterion.check()['score'] == criterion.check()['score'] == 1


def test_literature_references_inside_actual_docx_table(docx_document):
    document = docx.Document()
    document.add_heading('Раздел', 2)
    document.add_table(1, 1).cell(0, 0).text = 'Источник [1].'
    document.add_heading('Список использованных источников', 2)
    document.add_paragraph('Первый источник')
    assert ReferencesToLiteratureCheck(docx_document(document)).check()['score'] == 1


def test_banned_literature_domain_and_instance_isolation(md_document):
    info = md_document('# Название\n\n## Список использованных источников\n\n1. https://sky.pro/article')
    original = list(BannedWordsInLiteratureCheck.banned_words)
    criterion = BannedWordsInLiteratureCheck(info, banned_words=['custom.example'])
    assert criterion.check()['score'] == 0
    assert BannedWordsInLiteratureCheck.banned_words == original
    assert 'custom.example' not in BannedWordsInLiteratureCheck(info).banned_words


def test_header_references_reports_violation(md_document):
    result = ReportHeadersReferencesCheck(md_document('# Название\n\n## Заголовок [1]\n\nТекст')).check()
    assert result['score'] == 0
    assert 'содержит ссылки' in result['verdict'][0]


def test_lr_constructor_and_nested_paragraph_counts(docx_document):
    document = docx.Document()
    document.add_heading('Цель работы', 3)
    document.add_paragraph('Содержимое цели работы.')
    info = docx_document(document, report_type='LR')
    ReportChapters(info).check()  # No regex-as-style constructor failure.
    document = docx.Document()
    for name, level in [('1 Глава', 2), ('1.1 Раздел', 3), ('1.1.1 Подраздел', 4)]:
        document.add_heading(name, level)
        document.add_paragraph('Текст.')
    info = docx_document(document)
    result = ReportParagraphsCountCheck(info, min_paragraphs_in_section=1, min_paragraphs_in_subsection=1).check()
    assert result['score'] == 1
    assert (
        ReportParagraphsCountCheck(
            info, skip_sections=['1.1.1'], min_paragraphs_in_section=1, min_paragraphs_in_subsection=1
        ).check()['score']
        == 1
    )


def test_short_sections_use_both_limits(md_document):
    info = md_document('# Название\n\n## Раздел\n\nКороткий текст.')
    assert ReportShortSectionsCheck(info, min_section_count=1, min_section_len=2).check()['score'] == 1
    assert ReportShortSectionsCheck(info, min_section_count=2, min_section_len=2).check()['score'] == 0
    assert ReportShortSectionsCheck(info, min_section_count=1, min_section_len=3).check()['score'] == 0


def test_abstract_missing_pages_does_not_pass(docx_document):
    assert ReportMaxSizeOfAbstractCheck(docx_document()).check()['score'] == 0
    assert ReportFirstPagesCheck(docx_document(report_type='LR')).check()['score'] == 0


def test_main_character_uses_real_python_docx_tables(docx_document):
    document = docx.Document()
    for _ in range(8):
        document.add_table(1, 1).cell(0, 0).text = 'Факультет | КТИ'
    result = ReportMainCharacterCheck(docx_document(document)).check()
    assert isinstance(result['score'], float)


def test_empty_task_blocks_coordinates_are_not_indices(docx_document):
    info = docx_document()
    info['file'].pdf_file.get_text_on_page = lambda: {
        2: 'ЗАДАНИЕ\nСтудент Иванов Группа 123\nРуководитель Петров\nТема работы Проверка\n«01» мая 2026 «02» июня 2026'
    }
    assert EmptyTaskPageCheck(info).check()['score'] == 1
    info['file'].pdf_file.get_text_on_page = lambda: {2: 'ЗАДАНИЕ\nСтудент ___ Группа ___'}
    assert EmptyTaskPageCheck(info).check()['score'] == 0


@pytest.mark.parametrize('criterion', [CompareGoalAndContentCheck, CompareTasksAndContentCheck])
def test_semantic_empty_body_is_failure_not_exception(md_document, criterion):
    info = md_document(
        '# Название\n\n## Введение\n\nЦель работы: анализ данных.\n\nЗадачи:\n\n'
        '1. Обзор\n2. Анализ\n3. Сравнение\n\nОбъект исследования: данные'
    )
    assert criterion(info).check()['score'] == 0


def test_goal_at_zero_and_zero_similarity(md_document):
    info = md_document('# Название\n\n## Введение\n\nЦель работы: анализ данных.\n\n## Раздел\n\nАнализ данных.')
    assert CompareGoalAndContentCheck(info).check()['score'] == 1
    info = md_document('# Название\n\n## Введение\n\nЦель работы: анализ данных.\n\n## Раздел\n\nКошка спит.')
    assert CompareGoalAndContentCheck(info).check()['score'] == 0


def test_md_decimal_and_abbreviations_no_fake_pages(md_document):
    info = md_document('# Название\n\n## Раздел\n\nЧисло 1234.5678.')
    assert ReportDecimalPlacesCheck(info, max_violations=0).check()['score'] == 0
    result = ReportAbbreviationsCheck(info).check()
    assert 'Ошибка' not in result['verdict'][0]
