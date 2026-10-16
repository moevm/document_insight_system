"""Exercise a running deployment through its existing web and access_token API.

Run inside the web container (uses its configured secrets, never prints them):
    docker compose exec -T web python3 - < scripts/verify_deployment.py

Creates three labelled test checks in the existing database and leaves results
available for inspection. Uses the real broker, worker, conversions and storage.
"""

import hashlib
import io
import json
import os
import time
import traceback
import uuid
import zipfile
from pathlib import Path

import requests
from bson import ObjectId
from docx import Document

from app.build_info import get_build_info
from app.db.methods.client import get_db, get_fs

BASE = 'http://127.0.0.1:8080'
TOKEN = os.environ['ACCESS_TOKEN']
PREFIX = 'nonroot-e2e-' + uuid.uuid4().hex[:8]
REPORT = []


def record(kind, **kw):
    row = {'test': kind, **kw}
    REPORT.append(row)
    print(json.dumps(row, ensure_ascii=False), flush=True)


def ok(r, code=200):
    assert r.status_code == code, f'Expected HTTP {code}, got {r.status_code}: {r.text[:200]}'
    return r


def wait_result(check_id, session, task_id=None):
    end = time.monotonic() + 180
    while time.monotonic() < end:
        r = ok(requests.get(f'{BASE}/api/results/ready/{check_id}', timeout=10))
        if r.json().get('is_ended'):
            break
        time.sleep(1)
    else:
        raise AssertionError(f'Timed out waiting for {check_id}')
    check = get_db().checks.find_one({'_id': ObjectId(check_id)})
    assert check and check['is_ended'] and not check['is_failed'], str(check_id) + ' failed pipeline'
    assert isinstance(check['enabled_checks'], list) and len(check['enabled_checks']) >= 20
    valid = next(c for c in check['enabled_checks'] if c['id'] == 'simple_check')
    assert valid['score'] == 1, 'Parser failed: ' + str(valid)
    assert check['score'] >= 0
    errors = [
        c['id'] for c in check['enabled_checks'] if 'Информация об ошибке для администратора:' in str(c.get('verdict'))
    ]
    assert not errors, 'Internal criterion errors: ' + str(errors)
    page = ok(session.get(f'{BASE}/results/{check_id}', timeout=20))
    assert 'results_title' in page.text and 'results_table' in page.text
    svg = ok(requests.get(f'{BASE}/results/svg/{check_id}', timeout=10))
    assert 'Результат:' in svg.text and 'Работа проверяется' not in svg.text
    pdf = ok(session.get(f"{BASE}/get_pdf/{check['conv_pdf_fs_id']}", timeout=20))
    assert pdf.headers['Content-Type'].startswith('application/pdf') and pdf.content.startswith(b'%PDF')
    original = get_fs().open_download_stream(ObjectId(check_id)).read()
    converted = get_fs().open_download_stream(ObjectId(check['conv_pdf_fs_id'])).read()
    assert converted.startswith(b'%PDF')
    if task_id:
        task = ok(session.get(f'{BASE}/tasks/{task_id}', timeout=10)).json()
        assert task['task_status'] == 'SUCCESS', str(task['task_status'])
    logs = ok(session.get(f'{BASE}/api/results/logs/{check_id}', timeout=10)).json()
    assert logs['total'] > 0
    for suffix in (Path(check['filename']).suffix, '.pdf'):
        assert not Path('/usr/src/project/files', check_id + suffix).exists(), 'Worker did not clean files'
    record(
        'pipeline',
        id=check_id,
        filename=check['filename'],
        criteria=len(check['enabled_checks']),
        score=check['score'],
        criterion_errors=errors,
        pdf_bytes=len(pdf.content),
        source_bytes=len(original),
        logs=logs['total'],
    )
    return check


try:
    s = requests.Session()
    login = ok(
        s.post(
            BASE + '/login/',
            json={'username': 'admin', 'password_hash': hashlib.md5(os.environ['ADMIN_PASSWORD'].encode()).hexdigest()},
            timeout=20,
        )
    )
    assert login.text == 'admin', 'Admin login failed (configured password may differ from stored one)'
    for path, marker in [('/upload/', 'upload_file'), ('/check_list/', 'check_list'), ('/src/general.js', None)]:
        r = s.get(BASE + path, timeout=20)
        ok(r)
        if marker:
            assert marker in r.text
    record('web-login-and-pages', status='OK')
    metadata = get_build_info()
    version_page = ok(s.get(BASE + '/version/', timeout=20)).text
    for field in ('version', 'commit', 'branch', 'source_state'):
        assert metadata[field] in version_page, f'Build metadata {field} missing from /version/'
    record(
        'build-version', version=metadata['version'], commit=metadata['commit'], source_state=metadata['source_state']
    )
    doc = Document()
    doc.add_heading('Проверка развёртывания сервиса', 0)
    for title, text in [
        ('Реферат', 'Разработана система проверки отчётов.'),
        ('Содержание', 'Введение. Основная часть. Заключение.'),
        (
            'Введение',
            'Цель работы — проверить загрузку и обработку отчёта. Задачи: реализовать проверку; оценить результат.',
        ),
        ('Основная часть', 'Система сохраняет документы, выполняет проверки и показывает результат. ' * 30),
        ('Заключение', 'Была проверена загрузка документа, выполнена обработка и получен результат.'),
        ('Список использованных источников', '1. Документация проекта.'),
    ]:
        doc.add_page_break()
        doc.add_heading(title, 1)
        doc.add_paragraph(text)
    b = io.BytesIO()
    doc.save(b)
    body = b.getvalue()
    r = ok(
        s.post(
            BASE + '/upload/',
            files={
                'file': (
                    PREFIX + '-web.docx',
                    body,
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                )
            },
            data={'file_type': 'report'},
            timeout=30,
        )
    )
    ids = r.json()
    assert set(ids) == {'task_id', 'check_id'}
    check = wait_result(ids['check_id'], s, ids['task_id'])
    assert get_fs().open_download_stream(ObjectId(ids['check_id'])).read() == body
    record('web-upload-source-preserved', status='OK')
    md = (
        '# Проверка API\n\n## Введение\n\nЦель работы — проверить совместимость API.\n\n## Основная часть\n\n'
        + ('Была выполнена обработка документа и получен результат. ' * 35)
        + '\n\n## Заключение\n\nРабота выполнена.\n'
    ).encode()
    for full in (False, True):
        data = {'access_token': TOKEN, 'criteria': 'BaseReportCriterionPack'}
        if full:
            data['full_response'] = 'true'
        r = ok(
            requests.post(
                BASE + '/tasks/md',
                data=data,
                files={'file': (PREFIX + ('-api-json' if full else '-api-legacy') + '.md', md, 'text/markdown')},
                timeout=30,
            )
        )
        if full:
            ids = r.json()
            assert set(ids) == {'task_id', 'check_id'}
            check_id = ids['check_id']
            task_id = ids['task_id']
        else:
            assert ObjectId.is_valid(r.text.strip())
            check_id = r.text.strip()
            task_id = None
        check = wait_result(check_id, s, task_id)
        assert check['user'] == 'api_access_token'
        assert get_fs().open_download_stream(ObjectId(check_id)).read() == md
        record('api-response-contract', mode='JSON' if full else 'legacy-string', status='OK')
    for token in (None, 'deliberately-invalid'):
        data = {} if token is None else {'access_token': token}
        ok(
            requests.post(
                BASE + '/tasks/md', data=data, files={'file': ('invalid.md', md, 'text/markdown')}, timeout=20
            ),
            401,
        )
    record('api-rejects-missing-and-invalid-token', status='OK')
    for path in ('/get_csv/', '/get_zip/'):
        ok(requests.get(BASE + path, timeout=20), 403)
        r = ok(requests.get(BASE + path, params={'access_token': TOKEN, 'filter_filename': PREFIX}, timeout=30))
        if path == '/get_csv/':
            assert PREFIX in r.text
        else:
            with zipfile.ZipFile(io.BytesIO(r.content)) as z:
                assert any(n.endswith('.csv') for n in z.namelist())
                assert len(z.namelist()) >= 4
        record('api-token-export', endpoint=path, bytes=len(r.content), status='OK')
    Path('/tmp/dis-e2e-report.json').write_text(
        json.dumps({'prefix': PREFIX, 'checks': REPORT}, ensure_ascii=False, indent=2)
    )
    print('ALL DEPLOYMENT PIPELINE CHECKS PASSED', flush=True)
except Exception:
    print(traceback.format_exc().replace(TOKEN, '<redacted>'), flush=True)
    Path('/tmp/dis-e2e-report.json').write_text(
        json.dumps({'prefix': PREFIX, 'checks': REPORT, 'failed': True}, ensure_ascii=False, indent=2)
    )
    raise SystemExit(1) from None
